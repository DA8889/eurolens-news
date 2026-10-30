"""Export the last 24 hours of European articles from the GDELT Article List (GAL) to xlsx.

    python gdelt_gal.py                 # the latest 24 hours GDELT has published
    python gdelt_gal.py --hours 1       # shorter window for a quick test

The GAL runs about an hour behind real time (measured 2026-10-30: newest file 73 min
old), so the window ends at the newest published file rather than at "now".

The GAL is not a search API: GDELT publishes one JSON-lines file per minute (in practice
only a few minutes after each quarter hour have one) at
https://data.gdeltproject.org/gdeltv3/gal/YYYYMMDDHHMMSS.gal.json.gz, on a static file
server. "Last 24 hours" therefore means checking each of the 1,440 minutes and
downloading the files that exist. Each record has url, title, desc (one-sentence
description), domain, outletName, lang, date (publication), author, image.

The GAL has no source-country field, so the country comes from the outlet's domain:
its country ending (.de -> DE, .gr -> EL, .co.uk -> GB, .eu -> EU) or, for European
outlets on generic domains like .com, the KNOWN_OUTLETS table below.

Outputs (both git-ignored):
    exports/gdelt_gal_<start>_<hours>h.xlsx   sheets: Articles, Summary
    logs/gdelt_gal_<timestamp>.log            DEBUG log: every request, status, size, error
"""
import argparse
import gzip
import json
import logging
import sys
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

ROOT = Path(__file__).parent
BASE = "https://data.gdeltproject.org/gdeltv3/gal/"
USER_AGENT = "EuroLensNewsBot/0.1 (+https://github.com/DA8889/eurolens-news)"
WORKERS = 4          # polite: static files, but no need to hammer the server
RETRIES = 3

# Country endings of outlet domains -> EuroLens ISO-2 (Greece = EL, UK = GB).
COUNTRY_TLDS = {
    "at": "AT", "be": "BE", "bg": "BG", "cy": "CY", "cz": "CZ", "de": "DE", "dk": "DK",
    "ee": "EE", "gr": "EL", "es": "ES", "fi": "FI", "fr": "FR", "uk": "GB", "hr": "HR",
    "hu": "HU", "ie": "IE", "it": "IT", "lt": "LT", "lu": "LU", "lv": "LV", "mt": "MT",
    "nl": "NL", "pl": "PL", "pt": "PT", "ro": "RO", "se": "SE", "si": "SI", "sk": "SK",
    "eu": "EU",
}
# European outlets on generic domains (.com/.net/.org). Starter list; extend freely.
KNOWN_OUTLETS = {
    "dw.com": "DE", "euronews.com": "EU", "euractiv.com": "EU", "euobserver.com": "EU",
    "rapporteur.com": "EU", "france24.com": "FR", "bbc.com": "GB", "theguardian.com": "GB",
    "ft.com": "GB", "economist.com": "GB", "elpais.com": "ES", "ekathimerini.com": "EL",
    "irishtimes.com": "IE", "cyprus-mail.com": "CY", "timesofmalta.com": "MT",
    "lovinmalta.com": "MT", "novinite.com": "BG", "sofiaglobe.com": "BG",
    "romania-insider.com": "RO", "notesfrompoland.com": "PL", "brusselstimes.com": "BE",
}

log = logging.getLogger("gdelt_gal")


def setup_logging(log_path):
    log.setLevel(logging.DEBUG)
    fmt = logging.Formatter("%(asctime)s %(levelname)-7s %(message)s")
    fh = logging.FileHandler(log_path, encoding="utf-8")
    fh.setLevel(logging.DEBUG)
    fh.setFormatter(fmt)
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)
    ch.setFormatter(fmt)
    log.addHandler(fh)
    log.addHandler(ch)


def country_of(domain):
    """(ISO-2 or None, basis) for an outlet domain.

    >>> country_of("www.spiegel.de"), country_of("news.bbc.co.uk"), country_of("dw.com"), country_of("cnn.com")
    (('DE', 'tld'), ('GB', 'tld'), ('DE', 'known_outlet'), (None, None))
    """
    d = (domain or "").lower().removeprefix("www.")
    if d in KNOWN_OUTLETS:
        return KNOWN_OUTLETS[d], "known_outlet"
    tld = d.rsplit(".", 1)[-1]
    if tld in COUNTRY_TLDS:
        return COUNTRY_TLDS[tld], "tld"
    return None, None


def fetch_minute(session, minute):
    """Download one minute file. Returns (minute, records, status) and never raises."""
    name = minute.strftime("%Y%m%d%H%M00") + ".gal.json.gz"
    url = BASE + name
    for attempt in range(1, RETRIES + 1):
        t0 = time.monotonic()
        try:
            resp = session.get(url, timeout=60)
        except requests.RequestException as e:
            log.warning("%s attempt %d: %s: %s", name, attempt, type(e).__name__, e)
            time.sleep(2 * attempt)
            continue
        ms = round((time.monotonic() - t0) * 1000)
        if resp.status_code == 404:
            log.debug("%s 404 (no file this minute) %dms", name, ms)
            return minute, [], "missing"
        if resp.status_code in (429, 500, 502, 503, 504):
            log.warning("%s attempt %d: HTTP %d %dms", name, attempt, resp.status_code, ms)
            time.sleep(5 * attempt)
            continue
        if resp.status_code != 200:
            log.error("%s HTTP %d: %s", name, resp.status_code, resp.text[:200])
            return minute, [], f"http_{resp.status_code}"
        records, bad = [], 0
        try:
            text = gzip.decompress(resp.content).decode("utf-8")
        except (OSError, UnicodeDecodeError) as e:
            log.error("%s could not be decompressed (%d bytes): %s", name, len(resp.content), e)
            return minute, [], "corrupt"
        for n, line in enumerate(text.splitlines(), 1):
            if not line.strip():
                continue
            try:
                records.append(json.loads(line))
            except json.JSONDecodeError as e:
                bad += 1
                log.error("%s line %d is not valid JSON: %s", name, n, e)
        log.debug("%s 200 %d bytes %dms %d records%s", name, len(resp.content), ms,
                  len(records), f" ({bad} bad lines)" if bad else "")
        return minute, records, "ok"
    log.error("%s gave up after %d attempts", name, RETRIES)
    return minute, [], "failed"


def newest_file(session, now, max_hours=6):
    """The newest minute that has a GAL file, searching back from `now`; None if none."""
    for i in range(max_hours * 60):
        minute = now - timedelta(minutes=i)
        name = minute.strftime("%Y%m%d%H%M00") + ".gal.json.gz"
        try:
            status = session.head(BASE + name, timeout=30).status_code
        except requests.RequestException as e:
            log.warning("HEAD %s: %s", name, e)
            continue
        log.debug("HEAD %s %d", name, status)
        if status == 200:
            return minute
    return None


def clean(value):
    """Strip characters Excel cannot store; None stays None."""
    if value is None:
        return None
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--hours", type=float, default=24)
    a = ap.parse_args()

    run_ts = datetime.now(timezone.utc)
    (ROOT / "logs").mkdir(exist_ok=True)
    (ROOT / "exports").mkdir(exist_ok=True)
    log_path = ROOT / "logs" / f"gdelt_gal_{run_ts:%Y%m%dT%H%M%SZ}.log"
    setup_logging(log_path)

    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    latest = newest_file(session, run_ts.replace(second=0, microsecond=0))
    if latest is None:
        log.error("no GAL file in the last 6 hours; GDELT may be down (see %s)", log_path)
        sys.exit(1)
    log.info("newest GAL file is %s UTC, %d minutes behind real time",
             f"{latest:%H:%M}", (run_ts - latest).total_seconds() // 60)
    end = latest + timedelta(minutes=1)
    start = end - timedelta(hours=a.hours)
    minutes = [start + timedelta(minutes=i) for i in range(int((end - start).total_seconds() // 60))]
    log.info("window %s to %s UTC: checking %d minute files, %d at a time",
             f"{start:%Y-%m-%d %H:%M}", f"{end:%Y-%m-%d %H:%M}", len(minutes), WORKERS)

    t0 = time.monotonic()
    with ThreadPoolExecutor(WORKERS) as pool:
        results = list(pool.map(lambda m: fetch_minute(session, m), minutes))
    fetch_s = time.monotonic() - t0
    status = Counter(s for _, _, s in results)
    found = [m for m, _, s in results if s == "ok"]
    log.info("fetched in %.0fs: %s", fetch_s, dict(status))
    log.info("minutes past the quarter hour that had files: %s",
             dict(sorted(Counter(m.minute % 15 for m in found).items())))
    if not found:
        log.error("no GAL files found in the window; nothing to export (see %s)", log_path)
        sys.exit(1)

    seen_urls, total, rows, countries, basis = set(), 0, [], Counter(), Counter()
    for minute, records, _ in results:
        for r in records:
            total += 1
            url = r.get("url")
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
            cc, how = country_of(r.get("domain"))
            if not cc:
                continue
            countries[cc] += 1
            basis[how] += 1
            rows.append([r.get("date"), f"{minute:%Y-%m-%dT%H:%M}Z", cc, how,
                         r.get("outletName"), r.get("domain"), r.get("lang"),
                         r.get("title"), r.get("desc"), r.get("author"), url])
    with_desc = sum(1 for row in rows if row[8])
    log.info("records %d, unique urls %d, European %d (%d with a description)",
             total, len(seen_urls), len(rows), with_desc)

    out = ROOT / "exports" / f"gdelt_gal_{start:%Y%m%dT%H%M}Z_{a.hours:g}h.xlsx"
    t1 = time.monotonic()
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Articles")
    ws.append(["published", "seen_in_file", "source_country", "country_basis", "outlet",
               "domain", "language", "title", "description", "author", "url"])
    for row in rows:
        ws.append([clean(v) for v in row])
    sm = wb.create_sheet("Summary")
    for line in [
        ["window_start_utc", f"{start:%Y-%m-%d %H:%M}"], ["window_end_utc", f"{end:%Y-%m-%d %H:%M}"],
        ["minutes_checked", len(minutes)], ["files_found", status["ok"]],
        ["minutes_without_file", status["missing"]],
        ["files_failed", len(minutes) - status["ok"] - status["missing"]],
        ["fetch_seconds", round(fetch_s)], ["records_total", total], ["unique_urls", len(seen_urls)],
        ["european_articles", len(rows)], ["european_with_description", with_desc],
        ["country_from_tld", basis["tld"]], ["country_from_known_outlet", basis["known_outlet"]],
        [], ["source_country", "articles"], *sorted(countries.items(), key=lambda kv: -kv[1]),
        [], ["log_file", str(log_path.relative_to(ROOT))],
    ]:
        sm.append(line)
    wb.save(out)
    log.info("wrote %s (%.1f MB) in %.0fs", out.relative_to(ROOT),
             out.stat().st_size / 1e6, time.monotonic() - t1)
    log.info("log: %s", log_path.relative_to(ROOT))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        log.exception("crashed")
        raise
